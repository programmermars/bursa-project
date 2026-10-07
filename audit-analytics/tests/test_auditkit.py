import sys
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

from auditkit.benford import benford_test, first_two_digits  # noqa: E402
from auditkit.checks import CHECKS, duplicate_near, run_checks, split_purchases  # noqa: E402
from auditkit.cli import main  # noqa: E402
from auditkit.generate import Dataset, generate  # noqa: E402
from auditkit.scoring import evaluate  # noqa: E402


@pytest.fixture(scope="module")
def result():
    ds = generate(seed=11, n_invoices=4000, n_journals=6000)
    exc = run_checks(ds)
    return ds, exc, evaluate(ds, exc)


def test_generator_is_reproducible():
    a, b = generate(seed=3, n_invoices=1000, n_journals=500), generate(seed=3, n_invoices=1000, n_journals=500)
    pd.testing.assert_frame_equal(a.invoices, b.invoices)
    pd.testing.assert_frame_equal(a.truth, b.truth)


def test_every_planted_type_is_targeted_by_a_test(result):
    ds, _, _ = result
    targeted = {t for c in CHECKS for t in c.targets}
    assert set(ds.truth.anomaly) <= targeted


def test_rule_based_tests_find_what_they_target(result):
    _, _, ev = result
    rules = ev["by_test"][~ev["by_test"].test.str.contains("Isolation|Benford|unusual for account")]
    assert (rules.recall >= 0.95).all(), rules[rules.recall < 0.95]


def test_overall_coverage_and_ranking(result):
    _, _, ev = result
    cov = ev["coverage"]
    assert cov.found.sum() / cov.planted.sum() >= 0.9
    assert ev["precision_at"][25] >= 0.9  # records failing several tests are real problems


def test_legitimate_exceptions_lower_precision(result):
    """Year-end weekend closing and month-end late journals are flagged but are not planted problems."""
    _, _, ev = result
    bt = ev["by_test"].set_index("test")
    assert bt.loc["Posted on weekend / public holiday", "precision"] < 1
    assert bt.loc["Manual JE outside office hours", "precision"] < 1


def test_near_duplicate_ignores_monthly_rent():
    ds = generate(seed=5, n_invoices=1000, n_journals=500)
    rent = set(ds.invoices[ds.invoices.description == "Monthly rental / contract"].invoice_id)
    genuine = rent - set(ds.truth.record_id)  # planted duplicates may be copies of a rent invoice
    assert len(genuine) >= 60 and not set(duplicate_near(ds).record_id) & genuine


def test_split_purchase_rule():
    inv = pd.DataFrame({
        "invoice_id": ["A", "B", "C", "D"], "vendor_id": ["V1", "V1", "V1", "V2"], "estate": ["E1", "E1", "E2", "E1"],
        "invoice_date": pd.to_datetime(["2025-01-06", "2025-01-07", "2025-01-07", "2025-01-07"]),
        "amount": [3_000.0, 3_500.0, 4_000.0, 4_900.0]})
    ds = Dataset(*(pd.DataFrame() for _ in range(2)), inv, *(pd.DataFrame() for _ in range(4)))
    assert set(split_purchases(ds).record_id) == {"A", "B"}  # C is another estate, D another vendor


def test_benford_conformity():
    rng = np.random.default_rng(0)
    natural = pd.Series(10 ** rng.uniform(1, 6, 20_000))  # log-uniform data follows Benford exactly
    assert benford_test(natural).conclusion == "close conformity"
    made_up = pd.Series(rng.uniform(1_000, 9_999, 20_000))
    assert benford_test(made_up).conclusion == "nonconformity"
    assert list(first_two_digits(pd.Series([12.5, 987, 45_000, 5]))) == [12, 98, 45]


def test_cli_end_to_end_and_report(tmp_path):
    data, out = tmp_path / "data", tmp_path / "out"
    main(["generate", "--out", str(data), "--seed", "2"])
    main(["run", "--data", str(data), "--out", str(out)])
    wb = openpyxl.load_workbook(out / "exceptions_report.xlsx")
    assert {"Summary", "Risk ranking", "Accounts payable", "General ledger", "Benford", "Test scoring"} <= set(wb.sheetnames)
    assert "Status (Open / Explained / Issue)" in [c.value for c in wb["Risk ranking"][1]]
    assert "| Duplicate invoice (exact) |" in (out / "results.md").read_text()


def test_runs_on_real_data_without_answer_key(tmp_path):
    data, out = tmp_path / "data", tmp_path / "out"
    main(["generate", "--out", str(data), "--seed", "4"])
    (data / "truth.csv").unlink()
    main(["run", "--data", str(data), "--out", str(out)])
    wb = openpyxl.load_workbook(out / "exceptions_report.xlsx")
    assert "Test scoring" not in wb.sheetnames and not (out / "results.md").exists()


def test_same_signal_tests_are_not_double_counted():
    from auditkit.scoring import risk_ranking
    exc = pd.DataFrame({"table": ["journals"] * 3, "record_id": ["J1", "J1", "J2"],
                        "test": ["Benford", "Below limit", "Off hours"], "reason": ["", "", ""],
                        "weight": [1, 2, 2], "group": ["threshold", "threshold", "Off hours"]})
    r = risk_ranking(exc).set_index("record_id")
    assert r.loc["J1", "risk_score"] == 2 and r.loc["J2", "risk_score"] == 2
