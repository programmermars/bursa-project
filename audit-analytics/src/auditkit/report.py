"""Excel exception report, laid out like an audit working paper: summary, risk ranking, one sheet per area
with blank follow-up columns for the auditor, Benford table and test scoring."""
from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .benford import benford_test
from .generate import Dataset

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
FOLLOW_UP = ["Auditor", "Status (Open / Explained / Issue)", "Management explanation", "Evidence ref."]


def _format(ws, df: pd.DataFrame, widths: dict[str, int] | None = None) -> None:
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for i, col in enumerate(df.columns, start=1):
        sample = [len(str(col))] + [len(str(v)) for v in df[col].head(200)]
        ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(col, min(max(sample) + 2, 70))
    ws.freeze_panes = "A2"
    if len(df):
        ws.auto_filter.ref = ws.dimensions


def _sheet(writer, name: str, df: pd.DataFrame, **kw) -> None:
    df.to_excel(writer, sheet_name=name, index=False)
    _format(writer.sheets[name], df, **kw)


def write_report(ds: Dataset, exceptions: pd.DataFrame, ev: dict, path: Path, with_scoring: bool = True) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    ranking = ev["ranking"]
    bt = benford_test(ds.journals[(ds.journals.source == "manual")].amount)
    by_test = exceptions.groupby(["area", "test"]).record_id.nunique().rename("exceptions").reset_index()

    summary = pd.DataFrame([
        ("Population", f"{len(ds.invoices):,} invoices, {len(ds.journals):,} journal entries, {len(ds.vendors):,} vendors, "
                       f"{ds.estate_ops.estate.nunique()} estates"),
        ("Tests run", f"{exceptions.test.nunique()} tests with exceptions"),
        ("Records flagged", f"{len(ranking):,} unique records ({len(exceptions):,} exceptions)"),
        ("Highest-risk records", "See 'Risk ranking': records hit by several tests come first"),
        ("Benford (manual JEs, first two digits)", f"MAD {bt.mad:.4f}: {bt.conclusion} (n={bt.n:,})"),
    ] + ([("Precision of the top 25 / 50 / 100 ranked records",
           " / ".join(f"{v:.0%}" for v in ev["precision_at"].values()))] if with_scoring else []),
        columns=["Item", "Result"])

    with pd.ExcelWriter(path, engine="openpyxl") as w:
        _sheet(w, "Summary", summary, widths={"Item": 45, "Result": 100})
        start = len(summary) + 3
        by_test.to_excel(w, sheet_name="Summary", index=False, startrow=start)
        for cell in w.sheets["Summary"][start + 1]:
            cell.font = Font(bold=True)
        top = ranking.head(300).copy()
        for c in FOLLOW_UP:
            top[c] = ""
        _sheet(w, "Risk ranking", top, widths={"reasons": 110})
        for area, g in exceptions.groupby("area"):
            g = g.drop(columns=["area", "weight", "group"]).copy()
            for c in FOLLOW_UP:
                g[c] = ""
            _sheet(w, area[:31], g, widths={"reason": 90})
        bt_table = bt.table.reset_index(names="first_two_digits").round(4)
        _sheet(w, "Benford", bt_table)
        if with_scoring:
            _sheet(w, "Test scoring", ev["by_test"].round(3))
            _sheet(w, "Coverage", ev["coverage"].round(3))
    return path
